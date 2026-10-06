# -*- coding: utf-8 -*-
"""
Build College_BPUT_Affiliation_Check.xlsx
Cross-check contact-list colleges against official BPUT affiliated list (2022-23)
from https://www.bput.ac.in (Google Sheet linked as List of Affiliated Colleges - 2022-23).
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r"c:\xampp\htdocs\public_html")
SRC1 = ROOT / "Odisha_College_Contacts_Puri_Jagatsinghpur_Sundargarh.xlsx"
SRC2 = ROOT / "College_Contacts_Soro_Paralakhemundi_Bilaspur.xlsx"
BPUT_CSV = ROOT / "scripts" / "_bput_affiliated_2022_23.csv"
OUT = ROOT / "College_BPUT_Affiliation_Check.xlsx"

BPUT_SOURCE = (
    "BPUT official List of Affiliated Colleges 2022-23 "
    "(https://www.bput.ac.in → Quick Links / Google Sheet "
    "docs.google.com/spreadsheets/d/10dYEIBxKmSYQFH5GspqzsVCylfChE6er)"
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def norm(s: str) -> str:
    s = (s or "").lower()
    s = s.replace("&", " and ")
    s = re.sub(r"[^\w\s]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    # common abbreviations / noise
    for a, b in [
        ("mahavidyalaya", "college"),
        ("intitute", "institute"),
        ("institute of technology and management", "institute of technology management"),
        ("govt ", "government "),
        ("govt.", "government "),
        ("autonomous", ""),
        ("degree college", "college"),
        ("  ", " "),
    ]:
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def load_bput() -> list[dict]:
    rows = []
    with open(BPUT_CSV, encoding="utf-8", errors="replace") as f:
        for r in csv.DictReader(f):
            name = (r.get("NAME OF THE COLLEGE") or "").strip()
            if not name:
                continue
            rows.append(
                {
                    "code": (r.get("COL_CODE") or "").strip(),
                    "type": (r.get("TYPE") or "").strip(),
                    "name": name,
                    "courses": (r.get("COURSES OFFERED") or "").strip().replace("\\", "/"),
                    "district": (r.get("DISTRICT") or "").strip(),
                    "address": (r.get("ADDRESS") or "").strip(),
                    "norm": norm(name),
                }
            )
    return rows


def load_contact_colleges() -> list[dict]:
    from openpyxl import load_workbook

    colleges = []
    for path, state_hint in [
        (SRC1, "Odisha"),
        (SRC2, None),
    ]:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["All Colleges"]
        rows = list(ws.iter_rows(values_only=True))
        header_i = None
        for i, row in enumerate(rows):
            cells = [str(c or "") for c in row]
            if any(c.strip().lower() == "college name" for c in cells):
                header_i = i
                break
        if header_i is None:
            wb.close()
            continue
        header = [str(c or "").strip() for c in rows[header_i]]
        dist_i = next(i for i, h in enumerate(header) if "district" in h.lower())
        name_i = next(i for i, h in enumerate(header) if h.lower() == "college name")
        for row in rows[header_i + 1 :]:
            cells = [("" if c is None else str(c).strip()) for c in row]
            if len(cells) <= max(dist_i, name_i):
                continue
            loc = cells[dist_i]
            name = cells[name_i]
            if not name:
                continue
            colleges.append(
                {
                    "location": loc,
                    "name": name,
                    "source_file": path.name,
                    "state_hint": state_hint,
                }
            )
        wb.close()
    return colleges


def find_bput_match(name: str, location: str, bput: list[dict]) -> dict | None:
    n = norm(name)
    loc = (location or "").lower()

    # Strong exact / containment matches only (avoid false Yes)
    candidates = []
    for b in bput:
        bn = b["norm"]
        if n == bn or n in bn or bn in n:
            candidates.append((0, b))
            continue
        # token overlap for key distinctive phrases
        key_phrases = [
            "ghanashyam hemalata",
            "imt pharmacy",
            "college of pharmaceutical sciences puri",
            "sundargarh engineering college",
            "sundergarh engineering college",
            "balasore college of engineering",
            "satyasai engineering college",
            "srinix college",
            "modern engineering and management",
            "vijayanjali institute of technology",
            "academy of business administration",
            "cambridge school of management",
            "rourkela institute of technology",
            "rourkela institute of management",
            "kanak manjari",
            "iipm school of management",
            "dr ambedkar memorial institute of information",
        ]
        for kp in key_phrases:
            if kp in n and kp in bn:
                candidates.append((1, b))
                break

    if not candidates:
        return None

    # Prefer district alignment when available
    def dist_score(b):
        d = b["district"].lower().replace("sundergarh", "sundargarh")
        loc2 = loc.replace("soro", "balasore").replace("paralakhemundi", "gajapati")
        if d and d in loc2:
            return 0
        if loc2 in ("puri",) and d == "puri":
            return 0
        if loc2 in ("sundargarh",) and d == "sundargarh":
            return 0
        return 1

    candidates.sort(key=lambda x: (x[0], dist_score(x[1])))
    return candidates[0][1]


def classify_category(name: str, location: str) -> str:
    n = name.lower()
    if "bilaspur" in location.lower() or "chhattisgarh" in n:
        if any(x in n for x in ["polytechnic", "engineering college", "institute of engineering", "lcit", "j.k. institute", "chouksey engineering", "government engineering"]):
            if "polytechnic" in n or "girls polytechnic" in n:
                return "Polytechnic (CG)"
            return "Engg/Tech (CG)"
        if any(x in n for x in ["iti", "industrial training"]):
            return "ITI"
        if any(x in n for x in ["vishwavidyalaya", "university"]):
            return "University"
        if any(x in n for x in ["law"]):
            return "Law"
        if any(x in n for x in ["education", "shiksha", "iase", "teacher"]):
            return "Teacher Education"
        return "Degree (CG)"
    if any(x in n for x in ["iti", "industrial training", "itc"]):
        return "ITI"
    if "polytechnic" in n or "engineering school" in n or "school of engineering" in n:
        return "Polytechnic"
    if any(x in n for x in ["pharmacy", "pharmaceutical", "b.pharm", "bpharm"]):
        return "Pharmacy"
    if any(
        x in n
        for x in [
            "institute of technology",
            "engineering college",
            "college of engineering",
            "institute of engineering",
            "engineering & technology",
            "engineering and technology",
        ]
    ):
        return "Engg/Tech"
    if any(x in n for x in ["mba", "management and technology", "school of management", "institute of management"]):
        return "Management"
    if any(x in n for x in ["nursing"]):
        return "Nursing"
    if any(x in n for x in ["teacher education", "cte", "b.ed", "education"]):
        if "degree college of education" in n or "college of education & technology" in n or "college of education and technology" in n:
            return "Degree"
        if "teacher" in n or "cte" in n or "iase" in n or "shiksha" in n:
            return "Teacher Education"
    if any(x in n for x in ["cipet"]):
        return "Diploma/Tech (CIPET)"
    if any(x in n for x in ["+3", "degree", "mahavidyalaya", "women's college", "womens college", "autonomous college", "science college", "arts"]):
        return "Degree"
    if "university" in n:
        return "University"
    return "Other/Mixed"


def likely_body_and_status(name: str, location: str, bput_hit: dict | None) -> tuple[str, str, str, str]:
    """
    Returns: bput_flag (Yes/No/Unknown), likely_body, category, notes
    """
    loc = location.strip()
    loc_l = loc.lower()
    n = name.lower()
    category = classify_category(name, location)

    # Bilaspur CG — never BPUT
    if loc_l == "bilaspur":
        if "guru ghasidas" in n or "central university" in n:
            body = "Guru Ghasidas Vishwavidyalaya (Central University)"
        elif "atal bihari" in n or "abvv" in n or "bilaspur university" in n:
            body = "Atal Bihari Vajpayee Vishwavidyalaya (ABVV / Bilaspur University)"
        elif "pandit sundarlal" in n or "open" in n:
            body = "Pt. Sundarlal Sharma (Open) University / other CG body"
        elif "polytechnic" in n or "engineering college" in n or "institute of engineering" in n or "lcit" in n or "j.k. institute" in n or "chouksey engineering" in n or "government engineering" in n or "dr. c.v. raman institute of science and technology" in n or "mahamaya technical" in n:
            body = "CSVTU (Chhattisgarh Swami Vivekanand Technical University) / CG diploma body"
        else:
            body = "ABVV (Bilaspur University) / CSVTU / other Chhattisgarh body"
        return (
            "No",
            body,
            category,
            "Bilaspur is in Chhattisgarh — not under BPUT (Odisha). Marked Not BPUT.",
        )

    # Explicit BPUT match from official list
    if bput_hit:
        courses = bput_hit["courses"]
        body = f"BPUT (code {bput_hit['code']}; {bput_hit['type']})"
        note = f"Matched official BPUT affiliated list 2022-23: {bput_hit['name']} | Courses: {courses} | Dist: {bput_hit['district']}"
        # refine category from courses
        cl = courses.lower()
        if "pharm" in cl:
            category = "Pharmacy"
        elif "b.tech" in cl or "m.tech" in cl:
            category = "Engg/Tech"
        elif "mba" in cl or "mca" in cl:
            category = "Management/MCA"
        return "Yes", body, category, note

    # Hard No rules — Odisha
    if any(x in n for x in ["iti", "industrial training", "itc"]):
        return (
            "No",
            "NCVT / DGT (ITI)",
            "ITI",
            "ITIs are under NCVT/DGT skill training — not BPUT degree affiliation.",
        )

    if (
        "polytechnic" in n
        or "engineering school" in n
        or "school of engineering" in n
        or re.search(r"\bdiploma\b", n)
        or "ugie" in n
        or "utkalmani gopabandhu institute of engineering" in n
    ):
        return (
            "No",
            "SCTE&VT / DTE&T Odisha (Diploma)",
            "Polytechnic",
            "Diploma/polytechnic institutes are under SCTE&VT (exams) and DTE&T Odisha — not BPUT.",
        )

    if "cipet" in n:
        return (
            "No",
            "CIPET (Central Institute) / diploma training — not BPUT Balasore unit",
            "Diploma/Tech (CIPET)",
            "CIPET units are central institutes; BPUT list includes CIPET Bhubaneswar degree campus, not typically Balasore diploma unit. Not treated as BPUT here.",
        )

    if "centurion university" in n or "cutm" in n:
        return (
            "No",
            "Centurion University of Technology and Management (private university)",
            "University (Private)",
            "Former JITM was under BPUT until ~2010; now a private university awarding its own degrees — not BPUT-affiliated.",
        )

    if "nursing" in n:
        return (
            "No",
            "ONMC / health sciences university body (not BPUT)",
            "Nursing",
            "Nursing colleges are not under BPUT.",
        )

    if any(x in n for x in ["teacher education", "cte ", "cte)", "college of teacher education", "iase"]):
        return (
            "No",
            "State university / SCERT teacher-education framework",
            "Teacher Education",
            "Teacher education colleges are not BPUT technical affiliates.",
        )

    # Known diploma/tech names that look like engg but are SCTE&VT
    diploma_like = [
        "jhadeswar institute of engineering",
        "nilasaila institute of science",
        "ramarani institute of technology",
        "balasore school of engineering",
        "soro school of engineering",
        "satya sai school of engineering",
        "vijayanjali school of engineering",
        "b.i.t. polytechnic",
        "bit polytechnic",
        "odisha polytechnic",
        "asian workers development institute",
        "awdi",
    ]
    if any(x in n for x in diploma_like):
        extra = ""
        if "ramarani" in n:
            extra = " Website mentions possible future B.Tech under BPUT; current listed programmes are diploma (SCTE&VT) — marked Not BPUT for now."
        return (
            "No",
            "SCTE&VT / DTE&T Odisha (Diploma)",
            "Polytechnic",
            "Name suggests engineering/tech but institute runs diploma programmes under SCTE&VT, not BPUT B.Tech affiliation." + extra,
        )

    # Degree college heuristics by location
    degree_markers = [
        "degree",
        "mahavidyalaya",
        "women's college",
        "womens college",
        "autonomous college",
        "+3",
        "science college",
        "arts",
        "commerce",
        "municipal college",
        "ispat",
        "gandhi mahavidyalaya",
        "dalmia college",
        "fakir mohan",
        "s.k.c.g",
        "skcg",
        "svm autonomous",
        "s.c.s",
        "scs ",
        "college puri",
        "college soro",
        "college balasore",
        "college rourkela",
        "college sundargarh",
        "college jagatsinghpur",
        "college nilagiri",
        "college khaira",
        "college remuna",
        "college mitrapur",
        "college markona",
        "college oupada",
        "college kupari",
        "college avana",
        "college anantapur",
        "college pankapal",
        "college tirtol",
        "college paradip",
        "college paradeep",
        "college balikuda",
        "college kujang",
        "college bonaigarh",
        "college rajgangpur",
        "college hemgir",
        "college biramitrapur",
        "college birmitrapur",
        "college lephripara",
        "college lahunipara",
        "college koira",
        "college bargaon",
        "college subdega",
        "college mohana",
        "college gurandi",
        "college kashinagar",
        "college kasinagar",
        "college chandiput",
        "college ramagiri",
        "college khajuripada",
        "college sevakpur",
        "nimapara",
        "pipili college",
        "gop college",
        "balanga college",
        "konark bhagabati",
        "mangala mahavidyalaya",
        "neela saila",
        "vedvyas college",
        "rourkela college",
        "venus degree",
        "gayatri +3",
        "nirmaan degree",
        "c s degree",
        "cast bondamunda",
        "college of arts",
        "hrusikesh ray",
        "swarnachud",
        "upendranath",
        "u.n. college",
        "k.k.s. women",
        "h.k. mahatab",
        "belavoomi",
        "saraswata",
        "simulia college",
        "remuna degree",
        "nilgiri women",
        "nilagiri college",
        "khaira college",
        "oupada college",
        "soro women",
        "binodini science",
        "meena ketan",
        "hill top degree",
        "indira memorial",
        "parsuram degree",
        "sri ram degree",
        "sri venketeswar",
        "women's degree college, paralakhemundi",
        "baba saheb ambedkar",
        "dr. b.r. ambedkar national college",
        "baba panchalingeswar",
        "balikhand",
        "berhampur degree college, raj-berhampur",
        "jhadeswar +3",
        "sriram +3",
    ]

    is_degreeish = any(m in n for m in degree_markers) or (
        category == "Degree"
        and "institute of technology" not in n
        and "engineering" not in n
        and "pharmacy" not in n
        and "management and technology" not in n
        and "polytechnic" not in n
    )

    if is_degreeish:
        if loc_l == "puri" or loc_l == "jagatsinghpur":
            body = "Utkal University (typical for +3 arts/science/commerce in this region)"
        elif loc_l == "sundargarh":
            body = "Sambalpur University (typical for +3 colleges in Sundargarh/Rourkela region)"
        elif loc_l == "soro":
            body = "Fakir Mohan University (FMU), Balasore (typical for +3 colleges)"
        elif loc_l == "paralakhemundi":
            body = "Berhampur University (typical for +3 colleges in Gajapati)"
        else:
            body = "State university (Utkal / Sambalpur / FMU / Berhampur) — not BPUT"
        return (
            "No",
            body,
            "Degree",
            "General degree (+3) arts/science/commerce colleges affiliate to state universities, not BPUT.",
        )

    # Ambiguous tech/management names not on BPUT list
    if any(
        x in n
        for x in [
            "xavier institute of management and technology",
            "vision residential",
            "kalam institute of management",
            "mahalaxmi institute of technology",
            "raja kishore chandra academy",
            "bits, balasore",
            "bits balasore",
            "venus group",
            "nviron",
            "college of arts sc. tech",
            "cast ",
        ]
    ):
        return (
            "Unknown",
            "Verify — not found on BPUT 2022-23 affiliated list under this exact name",
            category if category != "Other/Mixed" else "Verify",
            "Name sounds technical/management or is unclear. Not present as a clear match in BPUT affiliated colleges 2022-23. Verify on https://www.bput.ac.in / academic audit portal before treating as BPUT.",
        )

    # Default for leftover Odisha names
    if category in ("Engg/Tech", "Pharmacy", "Management", "Management/MCA"):
        return (
            "Unknown",
            "Verify against current BPUT affiliation portal",
            category,
            "Looks technical/management but no confident match in BPUT 2022-23 list. Do not assume BPUT without confirmation.",
        )

    # Fallback degree-ish by location
    if loc_l in ("puri", "jagatsinghpur"):
        body = "Likely Utkal University / other non-BPUT Odisha body"
    elif loc_l == "sundargarh":
        body = "Likely Sambalpur University / SCTE&VT / other non-BPUT body"
    elif loc_l == "soro":
        body = "Likely FM University / SCTE&VT / other non-BPUT body"
    elif loc_l == "paralakhemundi":
        body = "Likely Berhampur University / SCTE&VT / other non-BPUT body"
    else:
        body = "Non-BPUT (verify)"
    return (
        "No",
        body,
        category,
        "No match on BPUT affiliated list; treated as Not BPUT based on name/type heuristics.",
    )


# ---------------------------------------------------------------------------
# Build workbook
# ---------------------------------------------------------------------------

def style_header(ws, row, fills, fonts, thin):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = fills["header"]
        cell.font = fonts["header"]
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin


def autosize(ws, min_w=10, max_w=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            if cell.value:
                length = max(length, min(max_w, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = max(min_w, length)


def main():
    bput = load_bput()
    colleges = load_contact_colleges()

    results = []
    for c in colleges:
        hit = find_bput_match(c["name"], c["location"], bput)
        flag, body, category, notes = likely_body_and_status(c["name"], c["location"], hit)
        source = BPUT_SOURCE if flag == "Yes" else (
            "Name/type heuristics + public affiliation patterns; Bilaspur = CG bodies"
            if c["location"].lower() == "bilaspur"
            else "Not on BPUT 2022-23 list / public SCTE&VT-DTE&T / state university patterns; see Notes"
        )
        if flag == "Unknown":
            source = "Needs verification on bput.ac.in / college website / OJEE seat matrix"
        results.append(
            {
                "Location": c["location"],
                "College Name": c["name"],
                "BPUT Affiliated?": flag,
                "Likely Affiliation Body": body,
                "Category": category,
                "Notes": notes,
                "Source": source,
                "Source File": c["source_file"],
                "BPUT Code": hit["code"] if hit else "",
                "BPUT Listed Name": hit["name"] if hit else "",
                "BPUT Courses": hit["courses"] if hit else "",
            }
        )

    # Counts
    by_loc = defaultdict(lambda: {"Yes": 0, "No": 0, "Unknown": 0, "Total": 0})
    totals = {"Yes": 0, "No": 0, "Unknown": 0}
    for r in results:
        loc = r["Location"]
        flag = r["BPUT Affiliated?"]
        by_loc[loc][flag] += 1
        by_loc[loc]["Total"] += 1
        totals[flag] += 1

    wb = Workbook()

    # Colors
    fills = {
        "header": PatternFill("solid", fgColor="0F3D5E"),
        "title": PatternFill("solid", fgColor="0B6E4F"),
        "yes": PatternFill("solid", fgColor="C6EFCE"),
        "no": PatternFill("solid", fgColor="FCE4D6"),
        "unk": PatternFill("solid", fgColor="FFF2CC"),
        "card": PatternFill("solid", fgColor="E8F1F8"),
        "soft": PatternFill("solid", fgColor="F4F7FA"),
        "accent": PatternFill("solid", fgColor="1B7A6E"),
    }
    fonts = {
        "header": Font(name="Calibri", bold=True, color="FFFFFF", size=11),
        "title": Font(name="Calibri", bold=True, color="FFFFFF", size=18),
        "h2": Font(name="Calibri", bold=True, color="0F3D5E", size=14),
        "h3": Font(name="Calibri", bold=True, color="0B6E4F", size=12),
        "body": Font(name="Calibri", size=11),
        "bold": Font(name="Calibri", bold=True, size=11),
    }
    thin = Border(
        left=Side(style="thin", color="B0BEC5"),
        right=Side(style="thin", color="B0BEC5"),
        top=Side(style="thin", color="B0BEC5"),
        bottom=Side(style="thin", color="B0BEC5"),
    )

    # ---------- Dashboard ----------
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "BPUT Affiliation Check"
    ws["A1"].font = fonts["title"]
    ws["A1"].fill = fills["title"]
    ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 32

    ws["A2"] = (
        "Colleges from Puri / Jagatsinghpur / Sundargarh / Soro(Balasore) / Paralakhemundi(Gajapati) / Bilaspur(CG) "
        "checked against BPUT affiliated colleges list (2022-23 official sheet). "
        "BPUT = Biju Patnaik University of Technology, Odisha (B.Tech / M.Tech / B.Pharm / MBA / MCA / Architecture etc.)."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 48

    ws["A4"] = "Overall Counts"
    ws["A4"].font = fonts["h2"]
    ws["A5"] = "BPUT (Yes)"
    ws["B5"] = totals["Yes"]
    ws["A6"] = "Not BPUT (No)"
    ws["B6"] = totals["No"]
    ws["A7"] = "Unknown / Verify"
    ws["B7"] = totals["Unknown"]
    ws["A8"] = "Total Colleges Checked"
    ws["B8"] = sum(totals.values())
    for r, fill in [(5, fills["yes"]), (6, fills["no"]), (7, fills["unk"])]:
        ws[f"A{r}"].fill = fill
        ws[f"B{r}"].fill = fill
        ws[f"A{r}"].font = fonts["bold"]
        ws[f"B{r}"].font = fonts["bold"]

    ws["A10"] = "Counts by Location"
    ws["A10"].font = fonts["h2"]
    headers = ["Location", "BPUT Yes", "Not BPUT", "Unknown", "Total", "% BPUT", "State"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(12, i, h)
    style_header(ws, 12, fills, fonts, thin)

    loc_order = ["Puri", "Jagatsinghpur", "Sundargarh", "Soro", "Paralakhemundi", "Bilaspur"]
    # include any unexpected
    for loc in sorted(by_loc.keys()):
        if loc not in loc_order:
            loc_order.append(loc)

    state_map = {
        "Puri": "Odisha",
        "Jagatsinghpur": "Odisha",
        "Sundargarh": "Odisha",
        "Soro": "Odisha (Balasore)",
        "Paralakhemundi": "Odisha (Gajapati)",
        "Bilaspur": "Chhattisgarh (NOT BPUT)",
    }

    row = 13
    for loc in loc_order:
        d = by_loc[loc]
        total = d["Total"] or 1
        vals = [
            loc,
            d["Yes"],
            d["No"],
            d["Unknown"],
            d["Total"],
            round(100.0 * d["Yes"] / total, 1),
            state_map.get(loc, ""),
        ]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row, i, v)
            cell.border = thin
            cell.font = fonts["body"]
            if i == 2 and v:
                cell.fill = fills["yes"]
            if i == 3:
                cell.fill = fills["no"]
            if i == 4 and v:
                cell.fill = fills["unk"]
            if loc == "Bilaspur":
                cell.fill = PatternFill("solid", fgColor="F8CBAD")
        row += 1

    row += 2
    ws.cell(row, 1, "Quick guide").font = fonts["h2"]
    row += 1
    tips = [
        "BPUT covers technical programmes: B.Tech, M.Tech, B.Pharm/M.Pharm, MBA, MCA, Architecture, Planning, some M.Sc — not general +3 degree colleges.",
        "Arts / Science / Commerce (+3) degree colleges → usually Utkal / Sambalpur / Berhampur / FM University — Not BPUT.",
        "Polytechnics / 'School of Engineering' diploma → SCTE&VT + DTE&T Odisha — Not BPUT.",
        "ITIs / ITCs → NCVT / DGT — Not BPUT.",
        "Bilaspur (Chhattisgarh) → CSVTU / ABVV / GGU etc. — Never BPUT.",
        "Only mark Yes when college appears on BPUT affiliated list (or equivalent official evidence).",
        f"Primary reference used: {BPUT_SOURCE}",
        "Note: Contact lists do not include every BPUT college in these districts (e.g. Sundargarh Engineering College, Balasore BCET/Srinix/MEMS etc. may be absent from the source Excels).",
    ]
    for t in tips:
        ws.cell(row, 1, "• " + t)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1).alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 30
        row += 1

    autosize(ws)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["G"].width = 28

    # ---------- Affiliation Check ----------
    ws2 = wb.create_sheet("Affiliation Check")
    cols = [
        "Location",
        "College Name",
        "BPUT Affiliated?",
        "Likely Affiliation Body",
        "Category",
        "Notes",
        "Source",
        "Source File",
        "BPUT Code",
        "BPUT Listed Name",
        "BPUT Courses",
    ]
    ws2.append(cols)
    style_header(ws2, 1, fills, fonts, thin)
    ws2.auto_filter.ref = f"A1:K{len(results)+1}"
    ws2.freeze_panes = "A2"

    flag_fill = {"Yes": fills["yes"], "No": fills["no"], "Unknown": fills["unk"]}
    for r in results:
        ws2.append([r[c] for c in cols])
        rr = ws2.max_row
        for col in range(1, len(cols) + 1):
            cell = ws2.cell(rr, col)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = fonts["body"]
        ws2.cell(rr, 3).fill = flag_fill.get(r["BPUT Affiliated?"], fills["soft"])
        ws2.cell(rr, 3).font = fonts["bold"]
        if r["Location"].lower() == "bilaspur":
            ws2.cell(rr, 1).fill = PatternFill("solid", fgColor="F8CBAD")

    widths = [14, 42, 14, 40, 16, 48, 36, 28, 12, 36, 22]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.row_dimensions[1].height = 28

    # ---------- BPUT Only ----------
    ws3 = wb.create_sheet("BPUT Only")
    ws3.append(cols)
    style_header(ws3, 1, fills, fonts, thin)
    yes_rows = [r for r in results if r["BPUT Affiliated?"] == "Yes"]
    for r in yes_rows:
        ws3.append([r[c] for c in cols])
        rr = ws3.max_row
        for col in range(1, len(cols) + 1):
            cell = ws3.cell(rr, col)
            cell.border = thin
            cell.fill = fills["yes"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if not yes_rows:
        ws3.append(["—", "No BPUT-affiliated colleges found in the source contact Excels", "", "", "", "", "", "", "", "", ""])
    for i, w in enumerate(widths, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"

    # ---------- Not BPUT ----------
    ws4 = wb.create_sheet("Not BPUT")
    ws4.append(cols)
    style_header(ws4, 1, fills, fonts, thin)
    for r in results:
        if r["BPUT Affiliated?"] != "No":
            continue
        ws4.append([r[c] for c in cols])
        rr = ws4.max_row
        for col in range(1, len(cols) + 1):
            cell = ws4.cell(rr, col)
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws4.cell(rr, 3).fill = fills["no"]
        if r["Location"].lower() == "bilaspur":
            ws4.cell(rr, 1).fill = PatternFill("solid", fgColor="F8CBAD")
    for i, w in enumerate(widths, 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.auto_filter.ref = f"A1:K{ws4.max_row}"
    ws4.freeze_panes = "A2"

    # Unknown sheet (helpful)
    wsU = wb.create_sheet("Unknown Verify")
    wsU.append(cols)
    style_header(wsU, 1, fills, fonts, thin)
    unk_rows = [r for r in results if r["BPUT Affiliated?"] == "Unknown"]
    for r in unk_rows:
        wsU.append([r[c] for c in cols])
        rr = wsU.max_row
        for col in range(1, len(cols) + 1):
            cell = wsU.cell(rr, col)
            cell.border = thin
            cell.fill = fills["unk"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if not unk_rows:
        wsU.append(["—", "No Unknown rows", "", "", "", "", "", "", "", "", ""])
    for i, w in enumerate(widths, 1):
        wsU.column_dimensions[get_column_letter(i)].width = w

    # ---------- Sources & Notes ----------
    ws5 = wb.create_sheet("Sources & Notes")
    ws5["A1"] = "Sources & Notes — What is BPUT?"
    ws5["A1"].font = fonts["title"]
    ws5["A1"].fill = fills["accent"]
    ws5.merge_cells("A1:B1")
    ws5.row_dimensions[1].height = 30

    content = [
        ("", ""),
        ("What BPUT is", ""),
        (
            "Explanation",
            "Biju Patnaik University of Technology (BPUT), Rourkela, is Odisha's technical university. "
            "It affiliates colleges that run AICTE-style technical programmes such as B.Tech, M.Tech, "
            "B.Pharm/M.Pharm, MBA, MCA, Architecture and Planning. It does NOT affiliate ordinary "
            "+3 Arts/Science/Commerce degree colleges.",
        ),
        ("", ""),
        ("What is NOT BPUT", ""),
        (
            "Degree arts/science/commerce",
            "Usually under state universities: Utkal University (coastal/central belts incl. Puri, Jagatsinghpur), "
            "Sambalpur University (western Odisha incl. Sundargarh), Fakir Mohan University (Balasore/Soro area), "
            "Berhampur University (southern Odisha incl. Gajapati/Paralakhemundi).",
        ),
        (
            "Polytechnics / diploma",
            "Under SCTE&VT (examining body) and DTE&T Odisha. Names like 'Polytechnic', 'Engineering School', "
            "or 'School of Engineering' are almost always diploma — Not BPUT.",
        ),
        (
            "ITIs / ITCs",
            "Under NCVT / DGT skill ecosystem — Not BPUT.",
        ),
        (
            "Bilaspur (Chhattisgarh)",
            "Outside Odisha. Engineering/tech typically CSVTU; general degree typically ABVV (Bilaspur University) "
            "or Guru Ghasidas Central University. Never BPUT.",
        ),
        (
            "Private universities",
            "Example: Centurion University (Paralakhemundi) — awards own degrees. Historically JITM was BPUT-affiliated "
            "until ~2010; today it is Not BPUT.",
        ),
        ("", ""),
        ("Sources used for this workbook", ""),
        (
            "Primary",
            "BPUT website https://www.bput.ac.in — Quick Link 'List of Affiliated Colleges - 2022-23' "
            "(Google Sheet id 10dYEIBxKmSYQFH5GspqzsVCylfChE6er). Also https://www.bput.ac.in/colleges-under-bput.html",
        ),
        (
            "Secondary",
            "College websites / Wikipedia / SCTE&VT-DTE&T public patterns for diploma vs degree distinction; "
            "Sundargarh Engineering College site (secsng.in) confirms BPUT for that college (not present in source Excels).",
        ),
        (
            "Caveat",
            "Affiliation can change year to year. The public consolidated sheet on bput.ac.in is labelled 2022-23. "
            "For 2024-26 confirmations use BPUT Affiliation / Academic Audit Portal (academicaudit.bput.ac.in) "
            "or the college's latest provisional affiliation letter. This workbook only marks Yes when matched "
            "to the official BPUT list.",
        ),
        ("", ""),
        ("BPUT colleges in these districts (from official list) — whether present in your Excels", ""),
        (
            "Puri (on BPUT list)",
            "Ghanashyam Hemalata Institute of Technology & Management; IMT Pharmacy College; "
            "College of Pharmaceutical Sciences, Puri — these three appear in your Puri contact Excel → Yes.",
        ),
        (
            "Sundargarh / Rourkela (on BPUT list)",
            "Sundargarh Engineering College; Rourkela Institute of Technology; DAMITS; IIPM School of Management; "
            "Kanak Manjari Institute of Pharmaceutical Sciences; Centre for UG & PG Studies (BPUT campus); "
            "Dr. Ambedkar Institute of Pharmaceutical Science — mostly ABSENT from your Sundargarh contact Excel "
            "(your list is heavy on +3 / polytechnic / ITI).",
        ),
        (
            "Balasore / Soro (on BPUT list)",
            "BCET; Satyasai Engineering College; Modern Engineering & Management Studies; Srinix; "
            "Vijayanjali Institute of Technology; ABA; RJ School of Management; Cambridge School of Management (Soro); "
            "several pharmacy/MBA institutes — mostly ABSENT from your Soro Excel (your list has many diploma "
            "'School of Engineering' / polytechnic names, which are Not BPUT).",
        ),
        (
            "Jagatsinghpur / Gajapati",
            "No colleges from these districts appear on the BPUT 2022-23 affiliated sheet. "
            "Paralakhemundi's Centurion University is Not BPUT.",
        ),
        ("", ""),
        ("Input files", ""),
        ("File 1", str(SRC1)),
        ("File 2", str(SRC2)),
        ("Output", str(OUT)),
        ("Generated", "Automated check — Yes only with BPUT-list evidence"),
    ]

    r = 2
    for k, v in content:
        if k and not v:
            ws5.cell(r, 1, k).font = fonts["h2"]
            ws5.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        else:
            ws5.cell(r, 1, k).font = fonts["bold"]
            ws5.cell(r, 2, v).alignment = Alignment(wrap_text=True)
            if k:
                ws5.cell(r, 1).fill = fills["card"]
            ws5.row_dimensions[r].height = max(30, 15 + 12 * (len(v) // 90))
        r += 1
    ws5.column_dimensions["A"].width = 34
    ws5.column_dimensions["B"].width = 100

    wb.save(OUT)
    print("Wrote", OUT)
    print("Totals", totals)
    print("By location:")
    for loc in loc_order:
        print(" ", loc, dict(by_loc[loc]))
    print("BPUT Yes colleges:")
    for r in yes_rows:
        print(" ", r["Location"], "|", r["College Name"], "|", r["BPUT Listed Name"])
    print("Unknown colleges:")
    for r in unk_rows:
        print(" ", r["Location"], "|", r["College Name"])


if __name__ == "__main__":
    main()
